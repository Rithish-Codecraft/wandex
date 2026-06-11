import os
import csv
import re
from typing import List, Dict, Any
import docx
from backend.config import settings

class MediaService:
    def __init__(self, chunk_size: int = 1500, chunk_overlap: int = 300):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def clean_text(self, text: str) -> str:
        # Replace multiple spaces/newlines
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    def parse_docx(self, file_path: str) -> str:
        """Extracts text paragraphs from a Word document."""
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also include table text
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    table_text.append(" | ".join(row_text))
        
        full_text = "\n\n".join(paragraphs)
        if table_text:
            full_text += "\n\n--- Tables ---\n" + "\n".join(table_text)
        return full_text

    def parse_csv(self, file_path: str) -> str:
        """Converts a CSV file into a clean textual/markdown representation."""
        lines = []
        with open(file_path, mode="r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for idx, row in enumerate(reader):
                if not any(row):
                    continue
                # Format header and rows
                if idx == 0:
                    lines.append(" | ".join(row))
                    lines.append("-+-".join(["---" for _ in row]))
                else:
                    lines.append(" | ".join(row))
        return "\n".join(lines)

    def parse_txt(self, file_path: str) -> str:
        """Reads raw text file."""
        with open(file_path, mode="r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def process_multimedia_file(self, file_path: str, mime_type: str) -> str:
        """
        Processes an image using OpenRouter multimodal payload,
        or returns a descriptive placeholder text for audio/video files.
        """
        import json
        import urllib.request
        import base64
        
        filename = os.path.basename(file_path)
        
        if mime_type.startswith("image/"):
            print(f"Processing image {filename} via OpenRouter multimodal completions...")
            try:
                with open(file_path, "rb") as f:
                    b64_data = base64.b64encode(f.read()).decode("utf-8")
                
                url = "https://openrouter.ai/api/v1/chat/completions"
                headers = {
                    "Authorization": f"Bearer {settings.OPENROUTER_API_KEY}",
                    "Content-Type": "application/json"
                }
                
                prompt = (
                    "Provide a detailed textual description of this image. Extract all readable text using OCR, "
                    "describe any diagrams, visual elements, charts, tables, or concepts depicted. "
                    "Ensure your output is structured and comprehensive so it can be indexed for retrieval."
                )
                
                model_name = settings.LLM_MODEL
                if not "/" in model_name:
                    model_name = f"google/{model_name}"
                    
                data = {
                    "model": model_name,
                    "messages": [
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": prompt},
                                {
                                    "type": "image_url",
                                    "image_url": {
                                        "url": f"data:{mime_type};base64,{b64_data}"
                                    }
                                }
                            ]
                        }
                    ]
                }
                
                req = urllib.request.Request(
                    url,
                    data=json.dumps(data).encode("utf-8"),
                    headers=headers,
                    method="POST"
                )
                
                with urllib.request.urlopen(req) as response:
                    res_body = json.loads(response.read().decode("utf-8"))
                    return res_body["choices"][0]["message"]["content"]
            except Exception as e:
                print(f"Failed to process image via OpenRouter: {e}")
                return f"[Image description error: {str(e)}]"
        else:
            # Audio/video transcription is not supported via OpenRouter
            file_type = "audio" if mime_type.startswith("audio/") else "video"
            return (
                f"[{file_type.capitalize()} content: filename {filename}]\n\n"
                f"Notice: Automatic {file_type} transcription is not available via OpenRouter. "
                f"Please provide a text or PDF version of the {file_type} file for indexing."
            )

    def chunk_text(self, text: str, source_name: str, title: str, author: str) -> List[Dict[str, Any]]:
        """Splits full text into chunks with overlap."""
        chunks = []
        cleaned = self.clean_text(text)
        
        if len(cleaned) <= self.chunk_size:
            chunks.append({
                "text": cleaned,
                "page": 1,
                "source": source_name,
                "title": title,
                "author": author,
                "chunk_id": f"{source_name}_p1_c0"
            })
            return chunks

        start = 0
        chunk_idx = 0
        while start < len(cleaned):
            end = start + self.chunk_size
            if end < len(cleaned):
                last_space = cleaned.rfind(' ', start, end)
                if last_space != -1 and last_space > start:
                    end = last_space
            
            chunk_text = cleaned[start:end].strip()
            if chunk_text:
                chunks.append({
                    "text": chunk_text,
                    "page": 1,
                    "source": source_name,
                    "title": title,
                    "author": author,
                    "chunk_id": f"{source_name}_p1_c{chunk_idx}"
                })
                chunk_idx += 1
            
            start = end - self.chunk_overlap
            if start >= len(cleaned) - self.chunk_overlap:
                break
                
        return chunks

    def parse_and_chunk(self, file_path: str, mime_type: str = None) -> List[Dict[str, Any]]:
        """
        Main entrypoint. Inspects file type, parses content, and splits into chunks.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        source_name = os.path.basename(file_path)
        title = source_name.rsplit(".", 1)[0]
        author = "Unknown"
        ext = source_name.split(".")[-1].lower()

        # Handle mapping extensions to MIME types if not provided
        if not mime_type:
            if ext in ["png", "jpg", "jpeg"]:
                mime_type = f"image/{ext}"
            elif ext in ["mp3", "wav"]:
                mime_type = f"audio/{ext}"
            elif ext in ["mp4", "avi"]:
                mime_type = f"video/{ext}"

        text_content = ""

        # Router based on file type
        if ext == "pdf":
            # PDFs are parsed page-by-page. For uniformity we fall back to PDFChunker if imported
            from rag.chunker import PDFChunker
            pdf_chunker = PDFChunker(self.chunk_size, self.chunk_overlap)
            return pdf_chunker.chunk_document(file_path)

        elif ext in ["docx", "doc"]:
            text_content = self.parse_docx(file_path)

        elif ext in ["csv"]:
            text_content = self.parse_csv(file_path)

        elif ext in ["txt", "md"]:
            text_content = self.parse_txt(file_path)

        elif mime_type and (mime_type.startswith("image/") or mime_type.startswith("audio/") or mime_type.startswith("video/")):
            text_content = self.process_multimedia_file(file_path, mime_type)

        else:
            # Fallback to reading as text if possible
            try:
                text_content = self.parse_txt(file_path)
            except Exception:
                raise ValueError(f"Unsupported file format: {ext}")

        if not text_content.strip():
            return []

        return self.chunk_text(text_content, source_name, title, author)

    def parse_url(self, url: str) -> str:
        """Fetches the webpage content and uses OpenRouter to extract clean main content."""
        import urllib.request
        from llm.gemini import generate_text
        try:
            req = urllib.request.Request(
                url, 
                headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
            )
            with urllib.request.urlopen(req, timeout=15) as response:
                html_bytes = response.read()
                html_text = html_bytes.decode('utf-8', errors='ignore')
            
            prompt = (
                f"You are a clean text extractor. Below is the raw HTML content of the webpage: {url}.\n"
                "Please extract the main article title, authors (if any), publication date (if any), and the main article body content.\n"
                "Discard all navigation menus, sidebar widgets, headers, footers, script/style code, and ads.\n"
                "Format the output as clean markdown text.\n\n"
                f"--- HTML CONTENT ---\n{html_text[:60000]}"
            )
            clean_markdown = generate_text(prompt)
            return clean_markdown
        except Exception as e:
            raise Exception(f"Failed to fetch or parse URL: {e}")

